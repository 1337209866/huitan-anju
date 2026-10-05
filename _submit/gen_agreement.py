# -*- coding: utf-8 -*-
"""人机协同履历表 v2 - 朱睿/刘浩威 正式版"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn

OUT = r"E:\Users\海之子\project\energy-carbon-manager\_submit\慧碳安居-人机协同履历表.docx"

doc = Document()
style = doc.styles['Normal']
style.font.name = '宋体'
style.font.size = Pt(10.5)
style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

for section in doc.sections:
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

def set_cn(run, name='宋体'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(18); r.bold = True
    set_cn(r, '黑体')

def subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(11)
    set_cn(r)

def h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.size = Pt(13); r.bold = True
    set_cn(r, '黑体')

def body(text, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.bold = bold
    set_cn(r)

def make_table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        p = hdr[i].paragraphs[0]
        rr = p.add_run(h)
        rr.bold = True; rr.font.size = Pt(10); set_cn(rr, '黑体')
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for row in rows:
        cells = t.add_row().cells
        for i, txt in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            rr = p.add_run(txt)
            rr.font.size = Pt(9.5)
            set_cn(rr)
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.TOP
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t


title('"慧碳安居"项目 · 人机协同履历表')
subtitle('参赛方向：④ 运营（保障房/产业园 · 能源管理 + ESG）　　项目负责人：朱睿')

h1('一、团队成员及分工')
make_table(
    ['姓名', '分工模块', '具体工作', '独立/AI辅助'],
    [
        ['朱睿\n（负责人）',
         '项目总体设计、异常检测、系统集成与技术文档',
         '负责项目需求分析与总体技术路线设计；明确建筑能耗监测、异常识别和能源管理的业务需求；设计系统总体架构及数据处理流程；分析建筑能耗数据特征，设计预测残差异常检测方法及两阶段鲁棒训练策略；负责异常检测模型参数设置、模型对比与结果分析；完成系统功能整合、运行测试、关键技术问题处理；负责项目技术文档、项目书及成果展示材料的撰写与修改。',
         '以人工独立完成为主。AI 仅用于辅助查询资料、代码局部补全、报错信息分析和文字表达优化，不参与核心方案决策及最终结果判断。'],
        ['刘浩威',
         '能耗数据模拟、储能调度、前端展示、碳核算与系统功能实现',
         '根据建筑实际用能规律设计模拟数据生成方法；建立办公、住宅、商业等不同建筑类型的典型负荷模型；加入温度影响、负荷波动、自相关等因素，提高模拟数据合理性；建立储能运行模型并利用 PuLP 完成储能优化调度；负责前端页面设计与数据可视化展示；完成碳排放核算功能、能源数据展示及系统交互功能；负责各功能模块联调、测试及运行结果检查。',
         '以人工独立设计和实现为主。AI 仅辅助 Python/JavaScript 代码补全、常见报错排查和部分技术资料检索，核心物理模型、调度逻辑、页面功能及参数均由成员自主确定。'],
    ],
    widths=[2.0, 3.3, 8.0, 3.5]
)

h1('二、使用的 AI 工具')
make_table(
    ['工具', '主要用途', '使用环节'],
    [
        ['豆包（Doubao）', '辅助技术资料查询、代码问题分析、程序报错解释、文字表达优化', '项目开发及文档整理过程中'],
        ['GitHub Copilot 类代码辅助工具', '辅助补全少量 Python、JavaScript 代码，提高重复性编码效率', '程序编写阶段'],
        ['搜索引擎', '查询建筑能耗标准、排放因子、能源价格及相关技术资料', '前期调研和参数核查阶段'],
    ],
    widths=[4.5, 8.5, 3.8]
)
body('说明：AI 工具主要承担辅助性、重复性工作，不直接决定项目技术路线、模型结构、核心参数和最终实验结论。')

h1('三、AI 参与的工作环节与人工验证')
make_table(
    ['环节', 'AI 主要提供的辅助', '人工完成的主要工作及验证方式', '最终成果归属'],
    [
        ['能耗数据生成',
         '根据需求提供负荷曲线和随机噪声模型的代码实现建议。',
         '人工根据建筑实际用能规律确定不同建筑类型的负荷特征；设置办公、住宅、商业等建筑的典型能耗水平；进一步加入温度影响、热惯性、AR(1) 自相关及天气变化等因素；通过多项指标检查生成数据是否符合实际规律。',
         '人工设计、人工验证，AI 仅辅助编码。'],
        ['异常检测算法',
         '提供 Isolation Forest、EWMA 等常见异常检测方法的实现思路和代码参考。',
         '人工分析不同算法在模拟建筑能耗数据上的适用性；对初始结果进行实际测试，发现 EWMA 误报较多、Isolation Forest 对渐变异常识别能力不足后，重新设计预测残差法，并结合两阶段鲁棒训练和 CUSUM 进行优化；人工完成参数调整、模型比较和最终性能评价。',
         '人工独立完成核心算法设计、优化与结果判断。'],
        ['储能优化调度',
         '辅助查询 PuLP 使用方法和补全部分程序代码。',
         '人工建立储能充放电约束、功率限制、容量约束及运行逻辑；根据建筑负荷和能源价格确定调度目标；人工检查优化结果是否满足储能运行规律，并对异常结果进行修正。',
         '人工完成模型建立、参数设置和结果验证。'],
        ['前端可视化',
         '辅助生成部分 ECharts 配置代码和 CSS 代码。',
         '人工确定页面结构、信息展示层级和功能布局；根据项目展示需求设计能耗监测、异常识别、储能调度及碳排放等页面；人工修改图表参数、交互逻辑和页面文字，并完成最终测试。',
         '人工设计和实现为主，AI 仅辅助重复性代码编写。'],
        ['碳排放核算',
         '辅助查询排放因子计算方法及代码表达。',
         '人工确定碳排放核算逻辑、数据来源和计算流程；根据能源消耗数据完成排放量计算，并对计算结果进行人工复核，确保单位、计算公式和结果数量级正确。',
         '人工独立完成核算逻辑和结果核查。'],
        ['系统联调与测试',
         '辅助分析程序报错和定位部分代码问题。',
         '人工完成不同模块之间的数据接口检查、运行流程测试和异常情况处理；重点检查异常检测结果、储能调度结果、碳排放数据及前端显示结果之间的一致性。',
         '人工完成系统测试和最终判断。'],
        ['技术文档',
         '对部分文字进行语句优化和格式建议。',
         '人工梳理项目研发过程、技术路线、实验结果和实际问题；根据真实开发过程重新组织文字，删除空泛表述，并对所有关键数据和实验结果进行回查。',
         '人工独立撰写和修改，AI 仅辅助语言表达。'],
    ],
    widths=[2.2, 3.8, 7.8, 3.5]
)

h1('四、最终成果中团队独立完成与 AI 辅助的边界')

body('团队独立完成的工作：', bold=True)
body('项目需求分析、应用场景确定、总体技术路线设计、系统功能划分、建筑能耗规律分析、模拟数据模型设计、异常检测方法选择与改进、模型参数设置、储能优化模型建立、储能运行约束设计、碳排放核算逻辑设计、前端页面结构设计、系统功能集成、程序测试、结果分析以及项目文档的最终撰写与修改，均由团队成员自主完成。')

body('朱睿主要负责：', bold=True)
body('项目总体方案设计、系统架构、异常检测算法、模型优化与参数调整、系统集成、技术路线梳理以及最终成果审核。针对异常检测过程中出现的误报和漏检问题，由成员自主分析原因并调整技术方案，而非直接采用 AI 推荐结果。')

body('刘浩威主要负责：', bold=True)
body('建筑能耗模拟数据生成、储能优化调度、前端功能实现、碳排放核算以及各模块之间的联调测试。数据模型中的建筑负荷特征、储能运行约束、碳核算方法和页面功能均由成员根据项目需求自主设计和验证。')

body('AI 主要承担辅助性工作：', bold=True)
body('包括少量 Python、JavaScript 代码补全、重复性代码编写、程序报错信息解释、技术资料检索以及部分文字表达优化。AI 生成的内容不会直接作为最终成果使用，所有涉及项目技术路线、模型参数、实验数据和结论的内容，均由团队成员进行判断、修改和验证。')

body('人工最终把关原则：', bold=True)
body('项目中的核心技术方案、关键参数和实验结果均经过人工审核。对于 AI 提出的方案，团队成员结合建筑能源管理实际需求和程序运行结果进行判断，合理的部分才予以参考，不合理的方案直接舍弃。最终成果以团队成员自主设计、实现、测试和验证的内容为准，AI 仅作为提高开发效率的辅助工具，不替代团队成员的专业判断和实际工作。')

doc.save(OUT)
print('saved:', OUT)
